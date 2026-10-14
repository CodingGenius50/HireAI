
from io import BytesIO

from pypdf import PdfReader
from docx import Document


def _read_file_bytes(file_obj):
    """Read a Django FileField or uploaded file safely."""
    file_obj.open("rb")

    try:
        return file_obj.read()
    finally:
        file_obj.close()


def extract_text_from_pdf(file_obj):
    """Extract selectable text from a PDF file."""
    file_bytes = _read_file_bytes(file_obj)
    reader = PdfReader(BytesIO(file_bytes))

    if reader.is_encrypted:
        raise ValueError("The PDF is password-protected.")

    pages = []

    for page in reader.pages:
        pages.append(page.extract_text() or "")

    text = "\n".join(pages).strip()

    if not text:
        raise ValueError(
            "No readable text found in PDF. "
            "The PDF may be scanned or image-only."
        )

    return text


def extract_text_from_docx(file_obj):
    """Extract text from DOCX paragraphs and tables."""
    file_bytes = _read_file_bytes(file_obj)
    document = Document(BytesIO(file_bytes))

    parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()

    if not text:
        raise ValueError("No readable text found in DOCX.")

    return text
